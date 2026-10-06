#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
HỆ SINH THÁI KHÉP KÍN BÓC TÁCH ĐỀ THI HÓA HỌC CHUẨN GDPT 2018
MODULE 4: docx_export_engine.py (Native Word OMML & GDPT 2018 DOCX Generator)
--------------------------------------------------------------------------------
Tính năng cốt lõi:
1. Chuyển đổi LaTeX Hóa học ($\mathrm{...}$) sang Word Native OMML (<m:oMath>).
2. Dàn trang chuẩn Bộ GD&ĐT (Nghị định 30/2020/NĐ-CP):
   - Căn lề: Top 2.0 cm, Bottom 2.0 cm, Left 3.0 cm, Right 2.0 cm.
   - Font: Times New Roman, dãn dòng 1.15 line, space after 3pt.
3. Phân chia chuẩn 3 phần cấu trúc khảo thí GDPT 2018:
   - Phần I: Trắc nghiệm 4 lựa chọn (Tự động dàn 1 hàng, 2 hàng hoặc 4 hàng).
   - Phần II: Trắc nghiệm Đúng/Sai (4 ý a-b-c-d).
   - Phần III: Trắc nghiệm trả lời ngắn.
4. Tự động giải mã và cắm ảnh từ mỏ neo [Slug_IMG_xx] vào đúng vị trí nội dung.
5. Hỗ trợ xuất 2 chế độ: Đề thi khảo thí độc lập hoặc Đề thi kèm Lời giải chi tiết.
6. 100% Python thuần, 0 Token AI, tương thích tuyệt đối với Microsoft Word.
================================================================================
"""


import io
import re
import base64
from typing import List, Dict, Any, Optional, Tuple
from xml.sax.saxutils import escape


import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn




# ==============================================================================
# 1. BỘ CHUYỂN ĐỔI LATEX HÓA HỌC SANG WORD NATIVE OMML
# ==============================================================================
def latex_to_omml(latex_str: str) -> str:
    r"""
    Chuyển đổi chuỗi LaTeX (đặc biệt là công thức hóa học \mathrm{...}) sang Word OMML XML.
    Hỗ trợ:
    - \mathrm{...} (chữ đứng chuẩn hóa học)
    - Subscript: _{2}, _2, _{12}
    - Superscript: ^{3+}, ^+, ^\circ
    - Sub-Superscript: _{a}^{b}
    - Fraction: \frac{a}{b}
    - Arrows & symbols: \rightarrow, \rightleftharpoons, \uparrow, \downarrow, \cdot, \equiv, \approx
    """
    clean = latex_str.strip()
    if clean.startswith('$') and clean.endswith('$'):
        clean = clean[1:-1].strip()


    # Bóc vỏ \mathrm{} nếu bao quanh toàn bộ
    m_mathrm = re.match(r'^\\mathrm\{(.+)\}$', clean, re.DOTALL)
    if m_mathrm:
        clean = m_mathrm.group(1).strip()


    # Thay thế các ký hiệu phản ứng và toán học đặc thù
    clean = clean.replace(r'\rightarrow', '→')
    clean = clean.replace(r'\rightleftharpoons', '⇌')
    clean = clean.replace(r'\uparrow', '↑')
    clean = clean.replace(r'\downarrow', '↓')
    clean = clean.replace(r'\equiv', '≡')
    clean = clean.replace(r'\cdot', '·')
    clean = clean.replace(r'\times', '×')
    clean = clean.replace(r'\div', '÷')
    clean = clean.replace(r'\pm', '±')
    clean = clean.replace(r'\approx', '≈')
    clean = clean.replace(r'\ne', '≠')
    clean = clean.replace(r'\neq', '≠')
    clean = clean.replace(r'\le', '≤')
    clean = clean.replace(r'\ge', '≥')
    clean = clean.replace(r'^\circ', '°')
    clean = clean.replace(r'\circ', '°')
    clean = clean.replace(r'\alpha', 'α')
    clean = clean.replace(r'\beta', 'β')
    clean = clean.replace(r'\gamma', 'γ')
    clean = clean.replace(r'\Delta', 'Δ')
    clean = clean.replace(r'\pi', 'π')
    clean = clean.replace(r'\sigma', 'σ')


    pos = 0
    length = len(clean)
    xml_parts = ['<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">']


    def make_run(text: str, is_math_style: bool = False) -> str:
        esc = escape(text)
        sty = '' if is_math_style else '<m:rPr><m:sty m:val="p"/></m:rPr>'
        return f'<m:r>{sty}<m:t>{esc}</m:t></m:r>'


    def extract_braced(s: str, start: int) -> Tuple[str, int]:
        if start >= len(s) or s[start] != '{':
            if start < len(s):
                return s[start], start + 1
            return '', start
        depth = 0
        end = start
        while end < len(s):
            if s[end] == '{':
                depth += 1
            elif s[end] == '}':
                depth -= 1
                if depth == 0:
                    return s[start + 1:end], end + 1
            end += 1
        return s[start + 1:], len(s)


    while pos < length:
        # Xử lý phân số \frac{num}{den}
        if clean.startswith(r'\frac', pos):
            pos += 5
            while pos < length and clean[pos] == ' ':
                pos += 1
            num_str, pos = extract_braced(clean, pos)
            while pos < length and clean[pos] == ' ':
                pos += 1
            den_str, pos = extract_braced(clean, pos)


            num_xml = make_run(num_str)
            den_xml = make_run(den_str)
            xml_parts.append(f'<m:f><m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>')
            continue


        ch = clean[pos]


        # Xử lý chỉ số dưới độc lập
        if ch == '_':
            pos += 1
            sub_str, pos = extract_braced(clean, pos)
            if pos < length and clean[pos] == '^':
                pos += 1
                sup_str, pos = extract_braced(clean, pos)
                xml_parts.append(
                    f'<m:sSubSup><m:e>{make_run("")}</m:e><m:sub>{make_run(sub_str)}</m:sub><m:sup>{make_run(sup_str)}</m:sup></m:sSubSup>'
                )
            else:
                xml_parts.append(f'<m:sSub><m:e>{make_run("")}</m:e><m:sub>{make_run(sub_str)}</m:sub></m:sSub>')
            continue
        elif ch == '^':
            pos += 1
            sup_str, pos = extract_braced(clean, pos)
            xml_parts.append(f'<m:sSup><m:e>{make_run("")}</m:e><m:sup>{make_run(sup_str)}</m:sup></m:sSup>')
            continue


        # Đọc cụm ký tự làm cơ số
        base_match = re.match(r'([A-Za-z0-9\+\-\=→⇌↑↓≡·×÷±≈≠≤≥°αβγΔπσ]+|\s+|\S)', clean[pos:])
        if base_match:
            base_token = base_match.group(1)
            pos += len(base_token)


            # Nếu cơ số có nhiều hơn 1 ký tự và theo sau là _ hoặc ^, tách ký tự cuối làm cơ số
            if len(base_token) > 1 and pos < length and (clean[pos] == '_' or clean[pos] == '^'):
                pre_text = base_token[:-1]
                xml_parts.append(make_run(pre_text))
                base_token = base_token[-1]


            has_sub = False
            has_sup = False
            sub_content = ""
            sup_content = ""


            if pos < length and clean[pos] == '_':
                has_sub = True
                pos += 1
                sub_content, pos = extract_braced(clean, pos)
                if pos < length and clean[pos] == '^':
                    has_sup = True
                    pos += 1
                    sup_content, pos = extract_braced(clean, pos)
            elif pos < length and clean[pos] == '^':
                has_sup = True
                pos += 1
                sup_content, pos = extract_braced(clean, pos)
                if pos < length and clean[pos] == '_':
                    has_sub = True
                    pos += 1
                    sub_content, pos = extract_braced(clean, pos)


            if has_sub and has_sup:
                xml_parts.append(
                    f'<m:sSubSup><m:e>{make_run(base_token)}</m:e><m:sub>{make_run(sub_content)}</m:sub><m:sup>{make_run(sup_content)}</m:sup></m:sSubSup>'
                )
            elif has_sub:
                xml_parts.append(f'<m:sSub><m:e>{make_run(base_token)}</m:e><m:sub>{make_run(sub_content)}</m:sub></m:sSub>')
            elif has_sup:
                xml_parts.append(f'<m:sSup><m:e>{make_run(base_token)}</m:e><m:sup>{make_run(sup_content)}</m:sup></m:sSup>')
            else:
                xml_parts.append(make_run(base_token))
        else:
            xml_parts.append(make_run(clean[pos]))
            pos += 1


    xml_parts.append('</m:oMath>')
    return "".join(xml_parts)




# ==============================================================================
# 2. XỬ LÝ ĐOẠN VĂN HỖN HỢP (TEXT, OMML & IMAGE ANCHOR)
# ==============================================================================
def add_mixed_content_to_paragraph(
    p: docx.text.paragraph.Paragraph,
    text: str,
    image_map: Optional[Dict[str, bytes]] = None,
    doc: Optional[docx.Document] = None,
    bold_prefix: Optional[str] = None
) -> None:
    """
    Phân tích và chèn nội dung hỗn hợp vào Paragraph:
    - Text thông thường (Times New Roman 12pt).
    - Công thức LaTeX $...$ -> OMML (<m:oMath>).
    - Mỏ neo [Slug_IMG_xx] -> Tải ảnh nhúng vào trang căn giữa.
    """
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)


    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True


    if not text:
        return


    # Tách token: Công thức $...$ hoặc Mỏ neo [Slug_IMG_xx]
    token_pattern = re.compile(r'(\$[^\$]+\$|\[[a-zA-Z0-9_]+_IMG_\d+\])')
    tokens = token_pattern.split(text)


    for token in tokens:
        if not token:
            continue


        # Công thức LaTeX
        if token.startswith('$') and token.endswith('$') and len(token) > 2:
            try:
                omml_str = latex_to_omml(token)
                omml_elm = parse_xml(omml_str)
                p._p.append(omml_elm)
            except Exception:
                r_fallback = p.add_run(token)
                r_fallback.font.name = 'Times New Roman'
                r_fallback.font.size = Pt(12)


        # Mỏ neo ảnh
        elif re.match(r'^\[[a-zA-Z0-9_]+_IMG_\d+\]$', token):
            anchor = token
            if image_map and anchor in image_map and doc is not None:
                img_bytes = image_map[anchor]
                try:
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(4)
                    img_p.paragraph_format.space_after = Pt(4)
                    r_img = img_p.add_run()
                    r_img.add_picture(io.BytesIO(img_bytes), width=Inches(3.2))


                    caption_p = doc.add_paragraph()
                    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption_p.paragraph_format.space_after = Pt(4)
                    r_cap = caption_p.add_run(f"Hình: {anchor}")
                    r_cap.font.name = 'Times New Roman'
                    r_cap.font.size = Pt(9.5)
                    r_cap.font.italic = True
                except Exception:
                    r_anc = p.add_run(f" {token} ")
                    r_anc.font.name = 'Times New Roman'
                    r_anc.font.size = Pt(11)
                    r_anc.font.bold = True
            else:
                r_anc = p.add_run(f" {token} ")
                r_anc.font.name = 'Times New Roman'
                r_anc.font.size = Pt(11)
                r_anc.font.bold = True


        # Text thuần
        else:
            r_txt = p.add_run(token)
            r_txt.font.name = 'Times New Roman'
            r_txt.font.size = Pt(12)




# ==============================================================================
# 3. ĐỘNG CƠ DỰNG ĐỀ THI DOCX CHUẨN GDPT 2018
# ==============================================================================
class ChemistryDocxExporter:
    """
    Động cơ xuất đề thi Hóa học GDPT 2018:
    - Thiết lập chuẩn khổ giấy A4, lề trang theo Nghị định 30.
    - Tạo tiêu đề đề thi quốc gia 2 cột không viền.
    - Dàn trang Phần I (A-D), Phần II (Đúng/Sai a-d), Phần III (Trả lời ngắn).
    - Cắm ảnh mỏ neo và công thức OMML không vỡ font.
    """


    def __init__(
        self,
        questions: List[Dict[str, Any]],
        images: Optional[Dict[str, Any]] = None,
        exam_title: str = "ĐỀ KIỂM TRA ĐÁNH GIÁ NĂNG LỰC HÓA HỌC",
        school_name: str = "SỞ GIÁO DỤC VÀ ĐÀO TẠO",
        academic_year: str = "NĂM HỌC 2026 - 2027",
        duration_minutes: int = 50,
        include_solutions: bool = False
    ):
        self.questions = questions
        self.raw_images = images or {}
        self.exam_title = exam_title
        self.school_name = school_name
        self.academic_year = academic_year
        self.duration_minutes = duration_minutes
        self.include_solutions = include_solutions


        # Chuẩn hóa image_map thành { anchor_tag: bytes }
        self.image_map: Dict[str, bytes] = {}
        self._prepare_image_map()


    def _prepare_image_map(self) -> None:
        """Chuyển đổi dữ liệu ảnh (Base64 hoặc byte thô) sang byte stream."""
        for key, val in self.raw_images.items():
            if isinstance(val, bytes):
                self.image_map[key] = val
            elif isinstance(val, str):
                try:
                    clean_b64 = val.split(',')[-1] if ',' in val else val
                    self.image_map[key] = base64.b64decode(clean_b64)
                except Exception:
                    pass
            elif isinstance(val, dict) and "base64_data" in val:
                try:
                    clean_b64 = val["base64_data"].split(',')[-1] if ',' in val["base64_data"] else val["base64_data"]
                    tag = val.get("anchor_tag", key)
                    self.image_map[tag] = base64.b64decode(clean_b64)
                except Exception:
                    pass


    def build_docx(self) -> docx.Document:
        """Tạo toàn bộ tài liệu Word hoàn chỉnh."""
        doc = docx.Document()


        # 1. Căn lề khổ A4 theo Nghị định 30/2020/NĐ-CP
        section = doc.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.79)     # 2.0 cm
        section.bottom_margin = Inches(0.79)  # 2.0 cm
        section.left_margin = Inches(1.18)    # 3.0 cm
        section.right_margin = Inches(0.79)   # 2.0 cm


        # 2. Tiêu đề hành chính chuẩn 2 cột không viền
        self._add_exam_header(doc)


        # 3. Phân loại câu hỏi thành 3 phần GDPT 2018
        part1_questions = []
        part2_questions = []
        part3_questions = []


        for q in self.questions:
            q_type = str(q.get("question_type", "")).upper()
            opts = q.get("options", [])


            if "ĐÚNG/SAI" in q_type or (opts and opts[0].get("label", "").lower() in ['a', 'b', 'c', 'd']):
                part2_questions.append(q)
            elif "TRẢ LỜI NGẮN" in q_type or not opts:
                part3_questions.append(q)
            else:
                part1_questions.append(q)


        # Nếu không tự nhận diện được, mặc định toàn bộ vào Phần I
        if not part2_questions and not part3_questions:
            part1_questions = self.questions


        # 4. Xuất từng phần
        if part1_questions:
            self._render_part_header(
                doc,
                "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
                f"Thí sinh trả lời từ câu {part1_questions[0].get('question_number', 1)} đến câu {part1_questions[-1].get('question_number', len(part1_questions))}. Mỗi câu hỏi thí sinh chỉ chọn một phương án."
            )
            for q in part1_questions:
                self._render_single_question(doc, q)


        if part2_questions:
            self._render_part_header(
                doc,
                "PHẦN II. Câu trắc nghiệm đúng sai.",
                f"Thí sinh trả lời từ câu {part2_questions[0].get('question_number', 1)} đến câu {part2_questions[-1].get('question_number', len(part2_questions))}. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai."
            )
            for q in part2_questions:
                self._render_true_false_question(doc, q)


        if part3_questions:
            self._render_part_header(
                doc,
                "PHẦN III. Câu trắc nghiệm trả lời ngắn.",
                f"Thí sinh trả lời từ câu {part3_questions[0].get('question_number', 1)} đến câu {part3_questions[-1].get('question_number', len(part3_questions))}."
            )
            for q in part3_questions:
                self._render_short_answer_question(doc, q)


        # 5. Chân trang kết thúc đề thi
        end_p = doc.add_paragraph()
        end_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        end_p.paragraph_format.space_before = Pt(14)
        r_end = end_p.add_run("---------- HẾT ----------")
        r_end.font.name = 'Times New Roman'
        r_end.font.size = Pt(11)
        r_end.font.bold = True


        note_p = doc.add_paragraph()
        note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_note = note_p.add_run("Thí sinh không được sử dụng tài liệu. Cán bộ coi thi không giải thích gì thêm.")
        r_note.font.name = 'Times New Roman'
        r_note.font.size = Pt(10)
        r_note.font.italic = True


        return doc


    def _add_exam_header(self, doc: docx.Document) -> None:
        """Tạo bảng tiêu đề đề thi chuẩn gồm 2 cột không đường viền."""
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False


        for row in table.rows:
            row.cells[0].width = Inches(3.0)
            row.cells[1].width = Inches(3.3)


        c_left = table.cell(0, 0)
        p_l1 = c_left.paragraphs[0]
        p_l1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_l1 = p_l1.add_run(self.school_name.upper())
        r_l1.font.name = 'Times New Roman'
        r_l1.font.size = Pt(10.5)


        p_l2 = c_left.add_paragraph()
        p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_l2 = p_l2.add_run("ĐỀ THI CHÍNH THỨC\n(Đề thi có các câu hỏi)")
        r_l2.font.name = 'Times New Roman'
        r_l2.font.size = Pt(10)
        r_l2.font.bold = True


        c_right = table.cell(0, 1)
        p_r1 = c_right.paragraphs[0]
        p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_r1 = p_r1.add_run(f"{self.exam_title.upper()}\n{self.academic_year.upper()}")
        r_r1.font.name = 'Times New Roman'
        r_r1.font.size = Pt(10.5)
        r_r1.font.bold = True


        p_r2 = c_right.add_paragraph()
        p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_r2 = p_r2.add_run(f"Môn thi: HÓA HỌC\nThời gian làm bài: {self.duration_minutes} phút (không kể phát đề)")
        r_r2.font.name = 'Times New Roman'
        r_r2.font.size = Pt(10)
        r_r2.font.italic = True


        self._remove_table_borders(table)


        sep_p = doc.add_paragraph()
        sep_p.paragraph_format.space_before = Pt(4)
        sep_p.paragraph_format.space_after = Pt(8)


    def _render_part_header(self, doc: docx.Document, title: str, instruction: str) -> None:
        """Tiêu đề từng phần của đề thi GDPT 2018."""
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(2)
        r_t = p_title.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(12)
        r_t.font.bold = True


        p_ins = doc.add_paragraph()
        p_ins.paragraph_format.space_after = Pt(6)
        r_i = p_ins.add_run(instruction)
        r_i.font.name = 'Times New Roman'
        r_i.font.size = Pt(11)
        r_i.font.italic = True


    def _render_single_question(self, doc: docx.Document, q: Dict[str, Any]) -> None:
        """Render một câu hỏi trắc nghiệm Phần I (4 phương án A-D)."""
        q_num = q.get("question_number", "")
        level = q.get("cognitive_level", "")
        level_str = f" [{level}]" if level else ""
        prefix = f"Câu {q_num}:{level_str} "


        p_stem = doc.add_paragraph()
        add_mixed_content_to_paragraph(
            p_stem,
            q.get("stem", ""),
            image_map=self.image_map,
            doc=doc,
            bold_prefix=prefix
        )


        options = q.get("options", [])
        if options:
            self._render_options_layout(doc, options)


        if self.include_solutions and q.get("solution") and q.get("solution").strip():
            self._render_solution_box(doc, q.get("solution").strip())


    def _render_true_false_question(self, doc: docx.Document, q: Dict[str, Any]) -> None:
        """Render một câu hỏi trắc nghiệm Đúng/Sai Phần II (4 ý a-d)."""
        q_num = q.get("question_number", "")
        level = q.get("cognitive_level", "")
        level_str = f" [{level}]" if level else ""
        prefix = f"Câu {q_num}:{level_str} "


        p_stem = doc.add_paragraph()
        add_mixed_content_to_paragraph(
            p_stem,
            q.get("stem", ""),
            image_map=self.image_map,
            doc=doc,
            bold_prefix=prefix
        )


        options = q.get("options", [])
        for opt in options:
            lbl = opt.get("label", "").lower()
            is_corr = opt.get("is_correct", False)
            corr_mark = " (Đúng)" if (self.include_solutions and is_corr) else ""
            opt_prefix = f"  {lbl}) "


            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Inches(0.25)
            add_mixed_content_to_paragraph(
                p_opt,
                opt.get("content", "") + corr_mark,
                image_map=self.image_map,
                doc=doc,
                bold_prefix=opt_prefix
            )


        if self.include_solutions and q.get("solution") and q.get("solution").strip():
            self._render_solution_box(doc, q.get("solution").strip())


    def _render_short_answer_question(self, doc: docx.Document, q: Dict[str, Any]) -> None:
        """Render một câu hỏi trả lời ngắn Phần III."""
        q_num = q.get("question_number", "")
        level = q.get("cognitive_level", "")
        level_str = f" [{level}]" if level else ""
        prefix = f"Câu {q_num}:{level_str} "


        p_stem = doc.add_paragraph()
        add_mixed_content_to_paragraph(
            p_stem,
            q.get("stem", ""),
            image_map=self.image_map,
            doc=doc,
            bold_prefix=prefix
        )


        if self.include_solutions and q.get("solution") and q.get("solution").strip():
            self._render_solution_box(doc, q.get("solution").strip(), label="Đáp số & Hướng dẫn giải")


    def _render_options_layout(self, doc: docx.Document, options: List[Dict[str, Any]]) -> None:
        """Tự động tính toán độ dài phương án để chọn bố cục 1 dòng, 2 dòng hay 4 dòng."""
        max_len = max([len(opt.get("content", "")) for opt in options]) if options else 0


        if len(options) == 4 and max_len <= 16:
            tbl = doc.add_table(rows=1, cols=4)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            self._remove_table_borders(tbl)
            for idx, opt in enumerate(options):
                lbl = opt.get("label", "").upper()
                is_c = opt.get("is_correct", False)
                prefix = f"*{lbl}. " if (self.include_solutions and is_c) else f"{lbl}. "
                cell = tbl.cell(0, idx)
                cell.width = Inches(1.55)
                p = cell.paragraphs[0]
                add_mixed_content_to_paragraph(p, opt.get("content", ""), image_map=self.image_map, doc=doc, bold_prefix=prefix)


        elif len(options) == 4 and max_len <= 45:
            tbl = doc.add_table(rows=2, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            self._remove_table_borders(tbl)
            for idx, opt in enumerate(options):
                row_idx = idx // 2
                col_idx = idx % 2
                lbl = opt.get("label", "").upper()
                is_c = opt.get("is_correct", False)
                prefix = f"*{lbl}. " if (self.include_solutions and is_c) else f"{lbl}. "
                cell = tbl.cell(row_idx, col_idx)
                cell.width = Inches(3.1)
                p = cell.paragraphs[0]
                add_mixed_content_to_paragraph(p, opt.get("content", ""), image_map=self.image_map, doc=doc, bold_prefix=prefix)


        else:
            for opt in options:
                lbl = opt.get("label", "").upper()
                is_c = opt.get("is_correct", False)
                prefix = f"*{lbl}. " if (self.include_solutions and is_c) else f"{lbl}. "
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                add_mixed_content_to_paragraph(p, opt.get("content", ""), image_map=self.image_map, doc=doc, bold_prefix=prefix)


    def _render_solution_box(self, doc: docx.Document, solution_text: str, label: str = "Lời giải chi tiết") -> None:
        """Tạo hộp lời giải chi tiết có viền xám nhẹ thanh lịch."""
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.3)


        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        tcPr.append(shd)


        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="12" w:space="0" w:color="3B82F6"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)


        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        add_mixed_content_to_paragraph(p, solution_text, image_map=self.image_map, doc=doc, bold_prefix=f"{label}: ")


        pad_p = doc.add_paragraph()
        pad_p.paragraph_format.space_after = Pt(4)


    def _remove_table_borders(self, table: docx.table.Table) -> None:
        """Gỡ bỏ đường viền của bảng OpenXML."""
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(tblBorders)




# ==============================================================================
# 4. GIAO DIỆN CÔNG CỘNG (PUBLIC API)
# ==============================================================================
def export_exam_to_docx_bytes(
    questions: List[Dict[str, Any]],
    images: Optional[Dict[str, Any]] = None,
    exam_title: str = "ĐỀ KIỂM TRA ĐÁNH GIÁ NĂNG LỰC HÓA HỌC",
    school_name: str = "SỞ GIÁO DỤC VÀ ĐÀO TẠO",
    academic_year: str = "NĂM HỌC 2026 - 2027",
    duration_minutes: int = 50,
    include_solutions: bool = False
) -> bytes:
    """Tạo đề thi và trả về luồng byte tệp DOCX."""
    exporter = ChemistryDocxExporter(
        questions=questions,
        images=images,
        exam_title=exam_title,
        school_name=school_name,
        academic_year=academic_year,
        duration_minutes=duration_minutes,
        include_solutions=include_solutions
    )
    doc = exporter.build_docx()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()




def export_exam_to_docx_base64(
    questions: List[Dict[str, Any]],
    images: Optional[Dict[str, Any]] = None,
    exam_title: str = "ĐỀ KIỂM TRA ĐÁNH GIÁ NĂNG LỰC HÓA HỌC",
    school_name: str = "SỞ GIÁO DỤC VÀ ĐÀO TẠO",
    academic_year: str = "NĂM HỌC 2026 - 2027",
    duration_minutes: int = 50,
    include_solutions: bool = False
) -> str:
    """Tạo đề thi và trả về chuỗi Base64."""
    docx_bytes = export_exam_to_docx_bytes(
        questions=questions,
        images=images,
        exam_title=exam_title,
        school_name=school_name,
        academic_year=academic_year,
        duration_minutes=duration_minutes,
        include_solutions=include_solutions
    )
    return base64.b64encode(docx_bytes).decode('utf-8')




# ==============================================================================
# 5. BỘ KIỂM THỬ THỰC NGHIỆM TÍCH HỢP (STRESS TEST)
# ==============================================================================
def run_export_stress_test() -> Dict[str, str]:
    """Kiểm thử nghiệm thu các bẫy biên khi dựng tệp Word DOCX."""
    report = {}


    mock_questions = [
        {
            "question_number": 1,
            "cognitive_level": "NB",
            "question_type": "TRẮC NGHIỆM",
            "stem": r"Chất $\mathrm{H_{2}SO_{4}}$ loãng tác dụng với $\mathrm{BaCl_{2}}$ sinh ra kết tủa $\mathrm{BaSO_{4}}\downarrow$ màu trắng.",
            "options": [
                {"label": "A", "content": "Màu đỏ.", "is_correct": False},
                {"label": "B", "content": "Màu trắng.", "is_correct": True},
                {"label": "C", "content": "Màu vàng.", "is_correct": False},
                {"label": "D", "content": "Màu tím.", "is_correct": False}
            ],
            "solution": r"Phương trình phản ứng: $\mathrm{BaCl_{2} + H_{2}SO_{4} \rightarrow BaSO_{4}\downarrow + 2HCl}$."
        },
        {
            "question_number": 2,
            "cognitive_level": "VD",
            "question_type": "ĐÚNG/SAI",
            "stem": r"Cho thí nghiệm điều chế khí như hình vẽ [H12_IMG_01].",
            "options": [
                {"label": "a", "content": r"Khí sinh ra là $\mathrm{NH_{3}}$.", "is_correct": True},
                {"label": "b", "content": "Thu khí bằng phương pháp dời nước.", "is_correct": False},
                {"label": "c", "content": r"Nồng độ dung dịch là $\frac{14}{131}$.", "is_correct": True},
                {"label": "d", "content": "Bông tẩm dung dịch kiềm.", "is_correct": False}
            ],
            "solution": "Khí nhẹ hơn không khí nên thu bằng phương pháp đẩy không khí úp ngược bình."
        },
        {
            "question_number": 3,
            "cognitive_level": "VDC",
            "question_type": "TRẢ LỜI NGẮN",
            "stem": r"Hòa tan hoàn toàn m gam Fe trong axit dư thu được ion $\mathrm{Fe^{3+}}$. Tính giá trị m.",
            "options": [],
            "solution": "Giá trị m = 5,6 gam."
        }
    ]


    mock_png_b64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="


    try:
        docx_bytes = export_exam_to_docx_bytes(
            questions=mock_questions,
            images={"[H12_IMG_01]": mock_png_b64},
            include_solutions=True
        )
        report["EXPORT_BYTES_GENERATED"] = "PASS" if len(docx_bytes) > 5000 else "FAIL"


        verify_doc = docx.Document(io.BytesIO(docx_bytes))
        report["DOCX_PARSED_BACK"] = "PASS" if len(verify_doc.paragraphs) > 5 else "FAIL"


        doc_xml_str = "".join([p._p.xml for p in verify_doc.paragraphs])
        report["OMML_EMBEDDED"] = "PASS" if "oMath" in doc_xml_str else "FAIL"


        b64_str = export_exam_to_docx_base64(
            questions=mock_questions,
            images={"[H12_IMG_01]": mock_png_b64},
            include_solutions=False
        )
        report["BASE64_EXPORT_VALID"] = "PASS" if len(b64_str) > 7000 else "FAIL"


    except Exception as e:
        report["EXCEPTION"] = f"FAIL ({str(e)})"


    return report




if __name__ == "__main__":
    print("=== ĐANG CHẠY BỘ KIỂM THỬ MODULE 4 (DOCX_EXPORT_ENGINE) ===")
    results = run_export_stress_test()
    for test_name, status in results.items():
        print(f"[{status}] {test_name}")
    print("============================================================")
